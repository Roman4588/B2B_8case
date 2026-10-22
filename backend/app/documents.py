import json
import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from pydantic import BaseModel, Field, ValidationError, field_validator


class RunbookInput(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    condition: str = Field(min_length=3, max_length=30000)
    steps: list[str] = Field(min_length=1, max_length=100)
    cautions: str = Field(default="", max_length=10000)
    category: str = "application"

    @field_validator("title", "condition")
    @classmethod
    def strip_text(cls, value):
        if not value.strip():
            raise ValueError("Поле не должно быть пустым")
        return value.strip()

    @field_validator("steps")
    @classmethod
    def valid_steps(cls, value):
        if any(not step.strip() or len(step) > 10000 for step in value):
            raise ValueError("Шаги должны содержать текст длиной до 10000 символов")
        return [step.strip() for step in value]

    @field_validator("category")
    @classmethod
    def valid_category(cls, value):
        if value not in {"server", "network", "database", "application"}:
            raise ValueError("Неизвестная категория")
        return value


def extract(content: bytes, extension: str) -> str:
    try:
        if extension in {".txt", ".md", ".json"}:
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            from pypdf import PdfReader

            if not content.startswith(b"%PDF-"):
                raise ValueError("Файл не является PDF")
            reader = PdfReader(BytesIO(content))
            if len(reader.pages) > 2000:
                raise ValueError("Лимит — 2000 страниц")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            from docx import Document

            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError("Слишком большой распакованный DOCX")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Файл не является DOCX")
            doc = Document(BytesIO(content))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            )
        else:
            raise ValueError("Поддерживаются JSON, MD, TXT, PDF и DOCX")
        if not text.strip():
            raise ValueError("Нет текстового слоя. Для скана сначала выполните OCR")
        if len(text) > 3_000_000:
            raise ValueError("Документ содержит слишком много текста")
        return text
    except (ValueError, UnicodeDecodeError, BadZipFile) as exc:
        raise ValueError(str(exc)) from exc
    except Exception as exc:
        raise ValueError("Не удалось прочитать документ: он повреждён или защищён") from exc


def parse_runbooks(text: str, extension: str, filename: str = "Документ") -> list[dict]:
    try:
        if extension == ".json":
            data = json.loads(text)
            records = data if isinstance(data, list) else [data]
        else:
            records = []
            for index, block in enumerate(re.split(r"(?m)^#{1,3}\s+", text)):
                block = block.strip()
                if not block:
                    continue
                match = re.search(
                    r"(?:Если|If)\s+(.+?)\s*(?:→|->|⇒|\bthen\b)\s*(.+)",
                    block,
                    flags=re.IGNORECASE | re.DOTALL,
                )
                if match:
                    condition, action = match.groups()
                    steps = [
                        re.sub(r"^[-*\d.)\s]+", "", s).strip() for s in re.split(r";|\n", action) if s.strip()
                    ]
                    records.append(
                        dict(title=condition.strip()[:200], condition=condition.strip(), steps=steps)
                    )
                else:
                    # Narrative runbooks remain searchable, with their original text as instruction.
                    for offset in range(0, len(block), 6000):
                        part = block[offset : offset + 6000]
                        records.append(
                            dict(
                                title=f"{filename[:150]} · раздел {index + 1}.{offset // 6000 + 1}",
                                condition=part[:2000],
                                steps=[part],
                            )
                        )
        if not records or len(records) > 2000:
            raise ValueError("Документ должен содержать от 1 до 2000 инструкций")
        return [RunbookInput.model_validate(row).model_dump() for row in records]
    except (ValidationError, json.JSONDecodeError, TypeError) as exc:
        raise ValueError("Неверный формат: нужны title, condition и непустой массив steps") from exc
